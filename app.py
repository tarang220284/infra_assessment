"""
app.py — FastAPI backend for the multi-domain Cyber-Recovery Discovery tool.

Surfaces:
  - Authentication (PBKDF2-SHA256 password hashing, opaque bearer tokens)
  - RBAC (admin / facilitator / viewer)
  - Projects → Sessions → Answers (project = customer engagement, sessions auto-named per domain)
  - Audit log (admin-only)
  - Multi-domain canonical model (8 domains from data/seed.json)
  - DOCX / Markdown / JSON exports

No external secrets manager. Single SQLite file under data/sessions.db.
Default admin: username 'admin' / password 'Alpha#2026' — must be reset on first login.
"""

from __future__ import annotations

import base64
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import socket
import sqlite3
import sys
import threading
import time
import uuid
import webbrowser
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from fastapi import Depends, FastAPI, Header, HTTPException, Query, Request
from fastapi.responses import (
    FileResponse, HTMLResponse, JSONResponse, PlainTextResponse, Response,
)
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import uvicorn

from docx import Document
from docx.shared import Pt

ROOT      = Path(__file__).resolve().parent
DATA      = ROOT / "data"
SEED_PATH = DATA / "seed.json"
DB_PATH   = DATA / "sessions.db"
STATIC    = ROOT / "static"
EXPORTS   = ROOT / "exports"
EXPORTS.mkdir(exist_ok=True)

# ────────────────────────────── canonical model ──────────────────────────────

if not SEED_PATH.exists():
    print("seed.json missing — running ingest.py first…", file=sys.stderr)
    import ingest  # type: ignore
    ingest.main()

with SEED_PATH.open(encoding="utf-8") as f:
    MODEL: Dict[str, Any] = json.load(f)
DOMAINS: Dict[str, Dict[str, Any]] = MODEL["domains"]
DOMAIN_ORDER: List[str] = MODEL["domain_order"]

QUESTION_LOOKUP: Dict[Tuple[str, int], Tuple[Dict[str, Any], Dict[str, Any]]] = {}
for dkey, d in DOMAINS.items():
    for l in d["layers"]:
        for q in l["questions"]:
            QUESTION_LOOKUP[(dkey, q["id"])] = (l, q)
QUESTION_IDS_BY_DOMAIN = {dk: [q["id"] for l in d["layers"] for q in l["questions"]] for dk, d in DOMAINS.items()}

# ────────────────────────────── password utilities ───────────────────────────

PBKDF2_ITERS = 600_000

def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERS)
    return f"pbkdf2_sha256${PBKDF2_ITERS}${base64.b64encode(salt).decode()}${base64.b64encode(digest).decode()}"

def verify_password(password: str, stored: str) -> bool:
    try:
        scheme, iters_s, salt_b64, hash_b64 = stored.split("$", 3)
        if scheme != "pbkdf2_sha256":
            return False
        iters = int(iters_s)
        salt = base64.b64decode(salt_b64)
        expected = base64.b64decode(hash_b64)
        candidate = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iters)
        return hmac.compare_digest(candidate, expected)
    except Exception:
        return False

PWD_RULE_RE = re.compile(r"^(?=.*[A-Z])(?=.*[0-9])(?=.*[^A-Za-z0-9]).{8,}$")
def validate_password(p: str) -> None:
    if not isinstance(p, str) or not PWD_RULE_RE.match(p):
        raise HTTPException(400, "Password must be ≥8 chars with at least 1 uppercase, 1 number, and 1 special character.")

# ────────────────────────────── sqlite layer ─────────────────────────────────

SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    id            TEXT PRIMARY KEY,
    username      TEXT UNIQUE NOT NULL COLLATE NOCASE,
    password_hash TEXT NOT NULL,
    role          TEXT NOT NULL DEFAULT 'facilitator',
    must_reset    INTEGER NOT NULL DEFAULT 0,
    active        INTEGER NOT NULL DEFAULT 1,
    full_name     TEXT DEFAULT '',
    created_at    TEXT NOT NULL,
    last_login_at TEXT
);
CREATE TABLE IF NOT EXISTS auth_tokens (
    token       TEXT PRIMARY KEY,
    user_id     TEXT NOT NULL,
    created_at  TEXT NOT NULL,
    expires_at  TEXT NOT NULL,
    FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS projects (
    id           TEXT PRIMARY KEY,
    name         TEXT NOT NULL,
    customer     TEXT NOT NULL,
    description  TEXT DEFAULT '',
    created_by   TEXT,
    created_at   TEXT NOT NULL,
    updated_at   TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS sessions (
    id            TEXT PRIMARY KEY,
    name          TEXT NOT NULL,
    domain        TEXT NOT NULL DEFAULT 'backup',
    customer      TEXT,
    facilitator   TEXT,
    project_id    TEXT,
    created_at    TEXT NOT NULL,
    updated_at    TEXT NOT NULL,
    FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS answers (
    session_id     TEXT NOT NULL,
    question_id    INTEGER NOT NULL,
    response       TEXT DEFAULT '',
    maturity       TEXT DEFAULT '',
    evidence       TEXT DEFAULT '',
    owner          TEXT DEFAULT '',
    status         TEXT DEFAULT 'unanswered',
    follow_up      TEXT DEFAULT 'none',
    follow_up_note TEXT DEFAULT '',
    not_applicable INTEGER NOT NULL DEFAULT 0,
    updated_at     TEXT NOT NULL,
    PRIMARY KEY (session_id, question_id),
    FOREIGN KEY (session_id) REFERENCES sessions(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS audit_logs (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id     TEXT,
    username    TEXT,
    action      TEXT NOT NULL,
    target_type TEXT DEFAULT '',
    target_id   TEXT DEFAULT '',
    detail      TEXT DEFAULT '',
    created_at  TEXT NOT NULL
);
"""

@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def _migrate(conn: sqlite3.Connection) -> None:
    scols = {r["name"] for r in conn.execute("PRAGMA table_info(sessions)")}
    if "project_id" not in scols:
        conn.execute("ALTER TABLE sessions ADD COLUMN project_id TEXT")
    if "domain" not in scols:
        conn.execute("ALTER TABLE sessions ADD COLUMN domain TEXT NOT NULL DEFAULT 'backup'")
    acols = {r["name"] for r in conn.execute("PRAGMA table_info(answers)")}
    if "not_applicable" not in acols:
        conn.execute("ALTER TABLE answers ADD COLUMN not_applicable INTEGER NOT NULL DEFAULT 0")
    # Backfill: orphan sessions get a 'Legacy {customer}' project
    rows = conn.execute("SELECT id, customer FROM sessions WHERE project_id IS NULL OR project_id = ''").fetchall()
    if rows:
        cust_to_proj: Dict[str, str] = {}
        t = now_iso()
        for r in rows:
            cust = (r["customer"] or "Unknown").strip() or "Unknown"
            if cust not in cust_to_proj:
                pid = uuid.uuid4().hex[:12]
                conn.execute(
                    "INSERT INTO projects(id, name, customer, description, created_by, created_at, updated_at) "
                    "VALUES(?,?,?,?,?,?,?)",
                    (pid, f"Legacy — {cust}", cust, "Auto-migrated from pre-project sessions", None, t, t),
                )
                cust_to_proj[cust] = pid
            conn.execute("UPDATE sessions SET project_id=? WHERE id=?", (cust_to_proj[cust], r["id"]))


def _seed_admin(conn: sqlite3.Connection) -> None:
    row = conn.execute("SELECT id FROM users WHERE role='admin' LIMIT 1").fetchone()
    if row:
        return
    uid = uuid.uuid4().hex[:12]
    conn.execute(
        "INSERT INTO users(id, username, password_hash, role, must_reset, active, full_name, created_at) "
        "VALUES(?,?,?,?,?,?,?,?)",
        (uid, "admin", hash_password("Alpha#2026"), "admin", 1, 1, "Default Admin", now_iso()),
    )
    print("Seeded default admin user — username='admin' password='Alpha#2026' (must reset on first login).",
          file=sys.stderr)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


with db() as _conn:
    _conn.executescript(SCHEMA)
    _migrate(_conn)
    _seed_admin(_conn)


def audit(conn: sqlite3.Connection, user: Optional[Dict[str, Any]], action: str,
          target_type: str = "", target_id: str = "", detail: str = "") -> None:
    conn.execute(
        "INSERT INTO audit_logs(user_id, username, action, target_type, target_id, detail, created_at) "
        "VALUES(?,?,?,?,?,?,?)",
        ((user or {}).get("id"), (user or {}).get("username", ""), action, target_type or "", target_id or "", detail or "", now_iso()),
    )

# ────────────────────────────── auth helpers ─────────────────────────────────

TOKEN_TTL_HOURS = 12

def _issue_token(conn: sqlite3.Connection, user_id: str) -> str:
    tok = secrets.token_urlsafe(32)
    now = datetime.now(timezone.utc)
    conn.execute(
        "INSERT INTO auth_tokens(token, user_id, created_at, expires_at) VALUES(?,?,?,?)",
        (tok, user_id, now.isoformat(timespec="seconds"), (now + timedelta(hours=TOKEN_TTL_HOURS)).isoformat(timespec="seconds")),
    )
    return tok


def _user_from_token(conn: sqlite3.Connection, token: str) -> Optional[Dict[str, Any]]:
    row = conn.execute(
        "SELECT u.* FROM users u JOIN auth_tokens t ON t.user_id = u.id "
        "WHERE t.token = ? AND t.expires_at > ? AND u.active = 1",
        (token, now_iso()),
    ).fetchone()
    return dict(row) if row else None


def get_current_user(authorization: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "Missing bearer token")
    token = authorization.split(" ", 1)[1].strip()
    with db() as conn:
        user = _user_from_token(conn, token)
        if not user:
            raise HTTPException(401, "Invalid or expired token")
        return user


def require_role(roles: List[str]):
    def dep(user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
        if user["role"] not in roles:
            raise HTTPException(403, "Forbidden")
        return user
    return dep

require_admin       = require_role(["admin"])
require_can_write   = require_role(["admin", "facilitator"])

def require_unblocked(user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    """User must NOT be in forced-reset mode for general API access."""
    if user.get("must_reset"):
        raise HTTPException(403, "Password reset required before accessing the system.")
    return user

# ────────────────────────────── pydantic models ──────────────────────────────

class LoginIn(BaseModel):
    username: str = Field(min_length=1, max_length=120)
    password: str = Field(min_length=1, max_length=200)

class PasswordResetIn(BaseModel):
    current_password: str
    new_password: str

class UserIn(BaseModel):
    username: str = Field(min_length=1, max_length=120)
    full_name: Optional[str] = ""
    role: str = Field(default="facilitator")
    initial_password: str

class UserPatch(BaseModel):
    full_name: Optional[str] = None
    role: Optional[str] = None
    active: Optional[bool] = None
    must_reset: Optional[bool] = None
    reset_password: Optional[str] = None

class ProjectIn(BaseModel):
    name: str = Field(min_length=1, max_length=160)
    customer: str = Field(min_length=1, max_length=160)
    description: Optional[str] = ""

class ProjectPatch(BaseModel):
    name: Optional[str] = None
    customer: Optional[str] = None
    description: Optional[str] = None

class SessionIn(BaseModel):
    project_id: str
    domain: str

class AnswerPatch(BaseModel):
    response:       Optional[str] = None
    maturity:       Optional[str] = None
    evidence:       Optional[str] = None
    owner:          Optional[str] = None
    status:         Optional[str] = None
    follow_up:      Optional[str] = None
    follow_up_note: Optional[str] = None
    not_applicable: Optional[bool] = None

VALID_ROLES = {"admin", "facilitator", "viewer"}

# ────────────────────────────── app ──────────────────────────────────────────

app = FastAPI(title="Infrastructure Assessment", docs_url=None, redoc_url=None)
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/", response_class=HTMLResponse, include_in_schema=False)
def index() -> FileResponse:
    return FileResponse(STATIC / "index.html")

# ────────────────── auth endpoints ───────────────────────────────────────────

@app.post("/api/auth/login")
def auth_login(payload: LoginIn) -> Dict[str, Any]:
    with db() as conn:
        row = conn.execute("SELECT * FROM users WHERE username = ? AND active = 1", (payload.username,)).fetchone()
        if not row or not verify_password(payload.password, row["password_hash"]):
            if row:
                audit(conn, {"id": row["id"], "username": row["username"]}, "login_failed",
                      target_type="user", target_id=row["id"])
            else:
                audit(conn, None, "login_failed", detail=f"unknown user: {payload.username[:80]}")
            raise HTTPException(401, "Invalid credentials")
        user = dict(row)
        token = _issue_token(conn, user["id"])
        conn.execute("UPDATE users SET last_login_at=? WHERE id=?", (now_iso(), user["id"]))
        audit(conn, user, "login", target_type="user", target_id=user["id"])
        return {
            "token": token,
            "user": _safe_user(user),
        }


@app.post("/api/auth/logout", status_code=204)
def auth_logout(authorization: Optional[str] = Header(default=None)) -> Response:
    if authorization and authorization.lower().startswith("bearer "):
        token = authorization.split(" ", 1)[1].strip()
        with db() as conn:
            user = _user_from_token(conn, token)
            conn.execute("DELETE FROM auth_tokens WHERE token = ?", (token,))
            if user:
                audit(conn, user, "logout", target_type="user", target_id=user["id"])
    return Response(status_code=204)


@app.get("/api/auth/me")
def auth_me(user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    return _safe_user(user)


@app.post("/api/auth/reset-password", status_code=204)
def auth_reset(payload: PasswordResetIn, user: Dict[str, Any] = Depends(get_current_user)) -> Response:
    if not verify_password(payload.current_password, user["password_hash"]):
        raise HTTPException(400, "Current password is incorrect.")
    validate_password(payload.new_password)
    if verify_password(payload.new_password, user["password_hash"]):
        raise HTTPException(400, "New password must differ from current.")
    with db() as conn:
        conn.execute("UPDATE users SET password_hash=?, must_reset=0 WHERE id=?",
                     (hash_password(payload.new_password), user["id"]))
        audit(conn, user, "password_reset", target_type="user", target_id=user["id"])
    return Response(status_code=204)


def _safe_user(u: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "id":            u["id"],
        "username":      u["username"],
        "full_name":     u.get("full_name") or "",
        "role":          u["role"],
        "must_reset":    bool(u.get("must_reset")),
        "active":        bool(u.get("active", 1)),
        "created_at":    u.get("created_at"),
        "last_login_at": u.get("last_login_at"),
    }

# ────────────────── canonical model ──────────────────────────────────────────

@app.get("/api/model")
def get_model(user: Dict[str, Any] = Depends(require_unblocked)) -> Dict[str, Any]:
    return MODEL

# ────────────────── projects ─────────────────────────────────────────────────

@app.get("/api/projects")
def list_projects(user: Dict[str, Any] = Depends(require_unblocked)) -> List[Dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT * FROM projects ORDER BY updated_at DESC").fetchall()
        out = []
        for r in rows:
            p = dict(r)
            counts = conn.execute(
                "SELECT domain, COUNT(*) AS n FROM sessions WHERE project_id=? GROUP BY domain",
                (p["id"],),
            ).fetchall()
            p["sessions_total"] = sum(c["n"] for c in counts)
            p["domains_assessed"] = len(counts)
            p["domains_total"] = len(DOMAIN_ORDER)
            out.append(p)
        return out


@app.post("/api/projects", status_code=201)
def create_project(payload: ProjectIn, user: Dict[str, Any] = Depends(require_can_write)) -> Dict[str, Any]:
    pid = uuid.uuid4().hex[:12]
    t = now_iso()
    with db() as conn:
        conn.execute(
            "INSERT INTO projects(id, name, customer, description, created_by, created_at, updated_at) VALUES(?,?,?,?,?,?,?)",
            (pid, payload.name.strip(), payload.customer.strip(), (payload.description or "").strip(), user["id"], t, t),
        )
        audit(conn, user, "project_create", target_type="project", target_id=pid,
              detail=f"name={payload.name!r} customer={payload.customer!r}")
        # Auto-create one session per domain — sessions are intrinsic to a project.
        _ensure_project_sessions(conn, pid, payload.customer.strip(), user)
        row = conn.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
        return dict(row)


@app.get("/api/projects/{pid}")
def get_project(pid: str, user: Dict[str, Any] = Depends(require_unblocked)) -> Dict[str, Any]:
    with db() as conn:
        p = conn.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
        if not p:
            raise HTTPException(404, "project not found")
        proj = dict(p)
        _ensure_project_sessions(conn, pid, proj["customer"], user)
        sess_rows = conn.execute("SELECT * FROM sessions WHERE project_id=? ORDER BY domain, updated_at DESC", (pid,)).fetchall()
        sessions = []
        for s in sess_rows:
            srow = dict(s)
            ensure_answer_rows(conn, srow["id"], srow["domain"])
            answers = fetch_answers(conn, srow["id"])
            srow["progress"] = _progress(srow["domain"], answers)
            sessions.append(srow)
        proj["sessions"] = sessions
        return proj


@app.patch("/api/projects/{pid}")
def update_project(pid: str, payload: ProjectPatch, user: Dict[str, Any] = Depends(require_can_write)) -> Dict[str, Any]:
    with db() as conn:
        cur = conn.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
        if not cur: raise HTTPException(404, "project not found")
        sets, vals = [], []
        for k, v in payload.dict(exclude_none=True).items():
            sets.append(f"{k}=?"); vals.append(v)
        if not sets: raise HTTPException(400, "no fields to update")
        sets.append("updated_at=?"); vals.append(now_iso())
        vals.append(pid)
        conn.execute(f"UPDATE projects SET {', '.join(sets)} WHERE id=?", vals)
        audit(conn, user, "project_update", target_type="project", target_id=pid, detail=json.dumps(payload.dict(exclude_none=True)))
        row = conn.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
        return dict(row)


@app.delete("/api/projects/{pid}", status_code=204)
def delete_project(pid: str, user: Dict[str, Any] = Depends(require_can_write)) -> Response:
    with db() as conn:
        cur = conn.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
        if not cur: raise HTTPException(404, "project not found")
        conn.execute("DELETE FROM projects WHERE id=?", (pid,))
        audit(conn, user, "project_delete", target_type="project", target_id=pid, detail=f"name={cur['name']!r}")
    return Response(status_code=204)


@app.get("/api/projects/{pid}/summary")
def project_summary(pid: str, user: Dict[str, Any] = Depends(require_unblocked)) -> Dict[str, Any]:
    with db() as conn:
        p = conn.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
        if not p: raise HTTPException(404, "project not found")
        _ensure_project_sessions(conn, pid, dict(p)["customer"], user)
        sess_rows = conn.execute(
            "SELECT * FROM sessions WHERE project_id=? ORDER BY updated_at DESC", (pid,)
        ).fetchall()
        per_domain: Dict[str, Dict[str, Any]] = {}
        for r in sess_rows:
            s = dict(r)
            ensure_answer_rows(conn, s["id"], s["domain"])
            answers = fetch_answers(conn, s["id"])
            summ = _summary(s, answers)
            if s["domain"] not in per_domain or s["updated_at"] > per_domain[s["domain"]]["updated_at"]:
                per_domain[s["domain"]] = {
                    "domain":       s["domain"],
                    "domain_label": DOMAINS[s["domain"]].get("label", s["domain"]),
                    "session_id":   s["id"],
                    "session_name": s["name"],
                    "updated_at":   s["updated_at"],
                    "overall_avg":  summ["overall_avg"],
                    "progress":     summ["progress"],
                    "by_layer":     summ["by_layer"],
                }
        ordered = [per_domain[d] for d in DOMAIN_ORDER if d in per_domain]
        scores = [x["overall_avg"] for x in ordered if x["overall_avg"] is not None]
        infra_avg = round(sum(scores)/len(scores), 2) if scores else None
        return {
            "project":          dict(p),
            "infra_avg":        infra_avg,
            "domains_assessed": len(ordered),
            "domains_total":    len(DOMAIN_ORDER),
            "per_domain":       ordered,
        }

# ────────────────── sessions ─────────────────────────────────────────────────

def ensure_answer_rows(conn: sqlite3.Connection, session_id: str, domain: str) -> None:
    existing = {r["question_id"] for r in conn.execute(
        "SELECT question_id FROM answers WHERE session_id=?", (session_id,))}
    missing = [qid for qid in QUESTION_IDS_BY_DOMAIN[domain] if qid not in existing]
    if not missing: return
    t = now_iso()
    conn.executemany(
        "INSERT INTO answers(session_id, question_id, updated_at) VALUES(?,?,?)",
        [(session_id, qid, t) for qid in missing],
    )


def fetch_session(conn: sqlite3.Connection, sid: str) -> Dict[str, Any]:
    row = conn.execute("SELECT * FROM sessions WHERE id=?", (sid,)).fetchone()
    if not row: raise HTTPException(404, f"session {sid} not found")
    return dict(row)


def fetch_answers(conn: sqlite3.Connection, sid: str) -> List[Dict[str, Any]]:
    rows = conn.execute("SELECT * FROM answers WHERE session_id=? ORDER BY question_id", (sid,)).fetchall()
    return [dict(r) for r in rows]


def _auto_session_name(domain: str) -> str:
    d = DOMAINS.get(domain)
    return (d.get("label") if d else domain) + " Discovery"


def _ensure_project_sessions(conn: sqlite3.Connection, project_id: str, customer: str,
                             actor: Optional[Dict[str, Any]]) -> int:
    """Idempotently create one session per domain for the given project. Returns count created."""
    existing = {r["domain"] for r in conn.execute(
        "SELECT domain FROM sessions WHERE project_id = ?", (project_id,)
    )}
    missing = [d for d in DOMAIN_ORDER if d not in existing]
    if not missing:
        return 0
    t = now_iso()
    for d in missing:
        sid = uuid.uuid4().hex[:12]
        conn.execute(
            "INSERT INTO sessions(id, name, domain, customer, facilitator, project_id, created_at, updated_at) "
            "VALUES(?,?,?,?,?,?,?,?)",
            (sid, _auto_session_name(d), d, customer, (actor or {}).get("username", "") or "", project_id, t, t),
        )
        ensure_answer_rows(conn, sid, d)
        audit(conn, actor, "session_auto_create", target_type="session", target_id=sid,
              detail=f"project={project_id} domain={d}")
    return len(missing)


@app.post("/api/sessions", status_code=201)
def create_session(payload: SessionIn, user: Dict[str, Any] = Depends(require_can_write)) -> Dict[str, Any]:
    if payload.domain not in DOMAINS:
        raise HTTPException(400, f"unknown domain '{payload.domain}'")
    with db() as conn:
        proj = conn.execute("SELECT * FROM projects WHERE id=?", (payload.project_id,)).fetchone()
        if not proj: raise HTTPException(400, "project not found")
        sid = uuid.uuid4().hex[:12]
        t = now_iso()
        name = _auto_session_name(payload.domain)
        conn.execute(
            "INSERT INTO sessions(id, name, domain, customer, facilitator, project_id, created_at, updated_at) "
            "VALUES(?,?,?,?,?,?,?,?)",
            (sid, name, payload.domain, proj["customer"], user["username"], payload.project_id, t, t),
        )
        ensure_answer_rows(conn, sid, payload.domain)
        audit(conn, user, "session_create", target_type="session", target_id=sid,
              detail=f"project={payload.project_id} domain={payload.domain}")
        return fetch_session(conn, sid)


@app.get("/api/sessions/{sid}")
def get_session(sid: str, user: Dict[str, Any] = Depends(require_unblocked)) -> Dict[str, Any]:
    with db() as conn:
        s = fetch_session(conn, sid)
        ensure_answer_rows(conn, sid, s["domain"])
        s["answers"] = fetch_answers(conn, sid)
        s["progress"] = _progress(s["domain"], s["answers"])
        return s


@app.delete("/api/sessions/{sid}", status_code=204)
def delete_session(sid: str, user: Dict[str, Any] = Depends(require_can_write)) -> Response:
    with db() as conn:
        cur = fetch_session(conn, sid)
        conn.execute("DELETE FROM sessions WHERE id=?", (sid,))
        audit(conn, user, "session_delete", target_type="session", target_id=sid, detail=f"name={cur['name']!r}")
    return Response(status_code=204)


@app.patch("/api/sessions/{sid}/answers/{qid}")
def patch_answer(sid: str, qid: int, patch: AnswerPatch, user: Dict[str, Any] = Depends(require_can_write)) -> Dict[str, Any]:
    with db() as conn:
        s = fetch_session(conn, sid)
        if (s["domain"], qid) not in QUESTION_LOOKUP:
            raise HTTPException(400, f"question {qid} not in domain {s['domain']}")
        ensure_answer_rows(conn, sid, s["domain"])
        sets, vals = [], []
        for k, v in patch.dict(exclude_none=True).items():
            if k == "not_applicable": v = 1 if v else 0
            sets.append(f"{k}=?"); vals.append(v)
        if not sets: raise HTTPException(400, "no fields to update")
        sets.append("updated_at=?"); vals.append(now_iso())
        vals.extend([sid, qid])
        conn.execute(f"UPDATE answers SET {', '.join(sets)} WHERE session_id=? AND question_id=?", vals)
        conn.execute("UPDATE sessions SET updated_at=? WHERE id=?", (now_iso(), sid))
        row = conn.execute("SELECT * FROM answers WHERE session_id=? AND question_id=?", (sid, qid)).fetchone()
        return dict(row)


@app.get("/api/sessions/{sid}/summary")
def session_summary(sid: str, user: Dict[str, Any] = Depends(require_unblocked)) -> Dict[str, Any]:
    with db() as conn:
        s = fetch_session(conn, sid)
        ensure_answer_rows(conn, sid, s["domain"])
        answers = fetch_answers(conn, sid)
    return _summary(s, answers)


@app.get("/api/sessions/{sid}/export/{fmt}")
def export_session(sid: str, fmt: str, user: Dict[str, Any] = Depends(require_unblocked)):
    fmt = fmt.lower()
    with db() as conn:
        s = fetch_session(conn, sid)
        ensure_answer_rows(conn, sid, s["domain"])
        answers = fetch_answers(conn, sid)
    fname = _safe_filename(f"{s['domain']}_{s['name']}") + "_" + datetime.now().strftime("%Y%m%d-%H%M%S")
    if fmt == "json":
        payload = {"session": s, "answers": answers, "summary": _summary(s, answers)}
        (EXPORTS / f"{fname}.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        return JSONResponse(content=payload, headers={"Content-Disposition": f'attachment; filename="{fname}.json"'})
    if fmt == "md":
        text = _render_markdown(s, answers, _summary(s, answers))
        (EXPORTS / f"{fname}.md").write_text(text, encoding="utf-8")
        return PlainTextResponse(text, headers={
            "Content-Disposition": f'attachment; filename="{fname}.md"',
            "Content-Type": "text/markdown; charset=utf-8"})
    if fmt == "docx":
        buf = _render_docx(s, answers, _summary(s, answers))
        (EXPORTS / f"{fname}.docx").write_bytes(buf.getvalue())
        return Response(content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{fname}.docx"'})
    if fmt == "summary-docx":
        buf = _render_summary_docx(s, answers, _summary(s, answers))
        (EXPORTS / f"{fname}_summary.docx").write_bytes(buf.getvalue())
        return Response(content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{fname}_summary.docx"'})
    raise HTTPException(400, f"unknown format '{fmt}'")

# ────────────────── admin endpoints ──────────────────────────────────────────

@app.get("/api/admin/users")
def admin_list_users(user: Dict[str, Any] = Depends(require_admin)) -> List[Dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT * FROM users ORDER BY created_at ASC").fetchall()
        return [_safe_user(dict(r)) for r in rows]


@app.post("/api/admin/users", status_code=201)
def admin_create_user(payload: UserIn, user: Dict[str, Any] = Depends(require_admin)) -> Dict[str, Any]:
    if payload.role not in VALID_ROLES:
        raise HTTPException(400, "invalid role")
    validate_password(payload.initial_password)
    with db() as conn:
        existing = conn.execute("SELECT id FROM users WHERE username = ?", (payload.username,)).fetchone()
        if existing:
            raise HTTPException(400, "username already exists")
        uid = uuid.uuid4().hex[:12]
        conn.execute(
            "INSERT INTO users(id, username, password_hash, role, must_reset, active, full_name, created_at) "
            "VALUES(?,?,?,?,?,?,?,?)",
            (uid, payload.username.strip(), hash_password(payload.initial_password), payload.role, 1, 1, (payload.full_name or "").strip(), now_iso()),
        )
        audit(conn, user, "user_create", target_type="user", target_id=uid,
              detail=f"username={payload.username!r} role={payload.role}")
        row = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        return _safe_user(dict(row))


@app.patch("/api/admin/users/{uid}")
def admin_update_user(uid: str, payload: UserPatch, user: Dict[str, Any] = Depends(require_admin)) -> Dict[str, Any]:
    with db() as conn:
        cur = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        if not cur: raise HTTPException(404, "user not found")
        sets, vals = [], []
        data = payload.dict(exclude_none=True)
        new_pw = data.pop("reset_password", None)
        if "role" in data and data["role"] not in VALID_ROLES:
            raise HTTPException(400, "invalid role")
        if "role" in data and data["role"] != "admin" and cur["role"] == "admin":
            # Block demoting the last admin
            n_admins = conn.execute("SELECT COUNT(*) AS n FROM users WHERE role='admin' AND active=1").fetchone()["n"]
            if n_admins <= 1: raise HTTPException(400, "Cannot demote the last admin.")
        for k, v in data.items():
            if k in ("active", "must_reset"): v = 1 if v else 0
            sets.append(f"{k}=?"); vals.append(v)
        if new_pw:
            validate_password(new_pw)
            sets.append("password_hash=?"); vals.append(hash_password(new_pw))
            sets.append("must_reset=?");    vals.append(1)
            # invalidate sessions
            conn.execute("DELETE FROM auth_tokens WHERE user_id=?", (uid,))
        if not sets: raise HTTPException(400, "no fields to update")
        vals.append(uid)
        conn.execute(f"UPDATE users SET {', '.join(sets)} WHERE id=?", vals)
        audit(conn, user, "user_update", target_type="user", target_id=uid,
              detail=json.dumps({k: ("***" if k == "reset_password" else v) for k, v in payload.dict(exclude_none=True).items()}))
        row = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        return _safe_user(dict(row))


@app.delete("/api/admin/users/{uid}", status_code=204)
def admin_delete_user(uid: str, user: Dict[str, Any] = Depends(require_admin)) -> Response:
    if uid == user["id"]:
        raise HTTPException(400, "cannot delete self")
    with db() as conn:
        cur = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        if not cur: raise HTTPException(404, "user not found")
        if cur["role"] == "admin":
            n = conn.execute("SELECT COUNT(*) AS n FROM users WHERE role='admin' AND active=1").fetchone()["n"]
            if n <= 1: raise HTTPException(400, "cannot delete the last admin")
        conn.execute("DELETE FROM users WHERE id=?", (uid,))
        audit(conn, user, "user_delete", target_type="user", target_id=uid, detail=f"username={cur['username']!r}")
    return Response(status_code=204)


@app.get("/api/admin/audit-logs")
def admin_audit_logs(limit: int = Query(default=200, ge=1, le=2000),
                     offset: int = Query(default=0, ge=0),
                     user: Dict[str, Any] = Depends(require_admin)) -> Dict[str, Any]:
    with db() as conn:
        total = conn.execute("SELECT COUNT(*) AS n FROM audit_logs").fetchone()["n"]
        rows = conn.execute(
            "SELECT * FROM audit_logs ORDER BY id DESC LIMIT ? OFFSET ?", (limit, offset),
        ).fetchall()
        return {"total": total, "limit": limit, "offset": offset, "items": [dict(r) for r in rows]}


# ────────────────── helpers (progress, summary, exports) ─────────────────────

def _progress(domain: str, answers: List[Dict[str, Any]]) -> Dict[str, Any]:
    d = DOMAINS[domain]
    by_layer: Dict[int, Dict[str, int]] = {
        l["id"]: {"total": len(l["questions"]), "complete": 0, "in_progress": 0, "not_applicable": 0}
        for l in d["layers"]
    }
    follow_ups = 0; complete = 0; na_count = 0
    for a in answers:
        key = (domain, a["question_id"])
        if key not in QUESTION_LOOKUP: continue
        layer = QUESTION_LOOKUP[key][0]["id"]
        is_na = bool(a.get("not_applicable"))
        if is_na:
            by_layer[layer]["not_applicable"] += 1
            by_layer[layer]["complete"] += 1
            complete += 1; na_count += 1
        else:
            if a["status"] == "complete":
                by_layer[layer]["complete"] += 1; complete += 1
            elif a["status"] == "in_progress":
                by_layer[layer]["in_progress"] += 1
        if a["follow_up"] == "open" and not is_na: follow_ups += 1
    total = len(answers) or 1
    return {
        "overall_pct": round(100 * complete / total, 1),
        "answers_complete": complete, "answers_total": total, "answers_na": na_count,
        "open_follow_ups": follow_ups, "by_layer": by_layer,
    }


def _summary(session: Dict[str, Any], answers: List[Dict[str, Any]]) -> Dict[str, Any]:
    domain = session["domain"]; d = DOMAINS[domain]
    score_map = {m["value"]: m["score"] for m in MODEL["maturity_scale"]}
    label_map = {m["value"]: m["label"] for m in MODEL["maturity_scale"]}
    heatmap = []
    layer_scores: Dict[int, List[int]] = {l["id"]: [] for l in d["layers"]}
    for a in answers:
        key = (domain, a["question_id"])
        if key not in QUESTION_LOOKUP: continue
        layer, q = QUESTION_LOOKUP[key]
        is_na = bool(a.get("not_applicable"))
        s = None if is_na else score_map.get(a["maturity"])
        if s is not None: layer_scores[layer["id"]].append(s)
        heatmap.append({
            "question_id": q["id"], "layer_id": layer["id"], "primary": q["primary"],
            "maturity": "" if is_na else a["maturity"],
            "maturity_label": "N/A" if is_na else label_map.get(a["maturity"], ""),
            "score": s, "status": a["status"], "follow_up": a["follow_up"],
            "not_applicable": is_na,
        })
    by_layer = []
    for l in d["layers"]:
        scores = layer_scores[l["id"]]
        avg = round(sum(scores)/len(scores), 2) if scores else None
        by_layer.append({"layer_id": l["id"], "title": l["title"], "avg_score": avg,
                         "answered": len(scores), "questions": len(l["questions"])})
    overall = [s for v in layer_scores.values() for s in v]
    return {
        "session_id": session["id"], "session_name": session["name"], "domain": domain,
        "domain_label": d.get("label", domain),
        "overall_avg": round(sum(overall)/len(overall), 2) if overall else None,
        "by_layer": by_layer, "heatmap": heatmap, "progress": _progress(domain, answers),
    }


def _safe_filename(s: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in s).strip("_") or "session"


def _render_markdown(session, answers, summary):
    domain = session["domain"]; d = DOMAINS[domain]
    a_by_qid = {a["question_id"]: a for a in answers}
    label_map = {m["value"]: m["label"] for m in MODEL["maturity_scale"]}
    out = []
    out += [f"# {d['title']} — Discovery Session", "",
            f"**Domain:** {d.get('label', domain)}  ",
            f"**Session:** {session['name']}  ",
            f"**Customer:** {session.get('customer') or '—'}  ",
            f"**Facilitator:** {session.get('facilitator') or '—'}  ",
            f"**Created:** {session['created_at']}  ",
            f"**Updated:** {session['updated_at']}  ", "",
            "## Summary",
            f"- Overall maturity (avg): **{summary['overall_avg']}/5**",
            f"- Progress: **{summary['progress']['answers_complete']}/{summary['progress']['answers_total']} complete** ({summary['progress']['overall_pct']}%)",
            f"- Open follow-ups: **{summary['progress']['open_follow_ups']}**",
            f"- Not-applicable answers: **{summary['progress']['answers_na']}**", "",
            "### Layer averages"]
    for row in summary["by_layer"]:
        out.append(f"- L{row['layer_id']} — {row['title']}: avg **{row['avg_score']}/5**, answered {row['answered']}/{row['questions']}")
    out += ["", f"> Facilitator intent: {d['facilitator_intro']}", ""]
    for l in d["layers"]:
        out += [f"## Layer {l['id']}: {l['title']}", f"*{l['intent']}*", ""]
        for q in l["questions"]:
            a = a_by_qid.get(q["id"], {})
            na = bool(a.get("not_applicable"))
            out += [f"### Q{q['id']}. {q['primary']}", "", "**Probing sub-questions**"]
            out += [f"- {sq}" for sq in q["sub_questions"]]
            out += ["", f"**Not Applicable:** {'Yes' if na else 'No'}  ",
                    f"**Maturity:** {'N/A' if na else label_map.get(a.get('maturity'), '—')}  ",
                    f"**Status:** {a.get('status') or 'unanswered'}  ",
                    f"**Owner:** {a.get('owner') or '—'}  ",
                    f"**Follow-up:** {a.get('follow_up') or 'none'}  ", "",
                    "**Response**", "", (a.get("response") or "_(no response captured)_").strip(), ""]
            if (a.get("evidence") or "").strip():
                out += ["**Evidence / notes**", "", a["evidence"].strip(), ""]
            if (a.get("follow_up_note") or "").strip():
                out += ["**Follow-up note**", "", a["follow_up_note"].strip(), ""]
    return "\n".join(out)


def _docx_styled(doc): style = doc.styles["Normal"]; style.font.name = "Arial"; style.font.size = Pt(10.5)


def _render_docx(session, answers, summary):
    domain = session["domain"]; d = DOMAINS[domain]
    doc = Document(); _docx_styled(doc)
    a_by_qid = {a["question_id"]: a for a in answers}
    label_map = {m["value"]: m["label"] for m in MODEL["maturity_scale"]}
    doc.add_heading(f"{d['title']} — Discovery Session", level=0)
    meta = doc.add_paragraph()
    meta.add_run("Domain: ").bold = True;      meta.add_run(d.get("label", domain) + "\n")
    meta.add_run("Session: ").bold = True;     meta.add_run(session["name"] + "\n")
    meta.add_run("Customer: ").bold = True;    meta.add_run((session.get("customer") or "—") + "\n")
    meta.add_run("Facilitator: ").bold = True; meta.add_run((session.get("facilitator") or "—") + "\n")
    meta.add_run("Updated: ").bold = True;     meta.add_run(session["updated_at"])
    doc.add_heading("Summary", level=1)
    p = doc.add_paragraph()
    p.add_run("Overall maturity (avg): ").bold = True; p.add_run(f"{summary['overall_avg']}/5\n")
    p.add_run("Progress: ").bold = True
    p.add_run(f"{summary['progress']['answers_complete']}/{summary['progress']['answers_total']} complete "
              f"({summary['progress']['overall_pct']}%)\n")
    p.add_run("Open follow-ups: ").bold = True; p.add_run(f"{summary['progress']['open_follow_ups']}\n")
    p.add_run("Not applicable: ").bold = True;  p.add_run(str(summary["progress"]["answers_na"]))
    t = doc.add_table(rows=1, cols=4); t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text = "Layer"; hdr[1].text = "Title"; hdr[2].text = "Avg / 5"; hdr[3].text = "Answered"
    for row in summary["by_layer"]:
        r = t.add_row().cells
        r[0].text = f"L{row['layer_id']}"; r[1].text = row["title"]
        r[2].text = "—" if row["avg_score"] is None else str(row["avg_score"])
        r[3].text = f"{row['answered']}/{row['questions']}"
    doc.add_paragraph().add_run("Facilitator intent: ").bold = True
    doc.add_paragraph(d["facilitator_intro"])
    for l in d["layers"]:
        doc.add_heading(f"Layer {l['id']}: {l['title']}", level=1)
        doc.add_paragraph().add_run(l["intent"]).italic = True
        for q in l["questions"]:
            a = a_by_qid.get(q["id"], {})
            na = bool(a.get("not_applicable"))
            doc.add_heading(f"Q{q['id']}. {q['primary']}", level=2)
            doc.add_paragraph().add_run("Probing sub-questions").bold = True
            for sq in q["sub_questions"]: doc.add_paragraph(sq, style="List Bullet")
            mp = doc.add_paragraph()
            mp.add_run("Not Applicable: ").bold = True; mp.add_run(("Yes" if na else "No") + "    ")
            mp.add_run("Maturity: ").bold = True; mp.add_run(("N/A" if na else label_map.get(a.get("maturity"), "—")) + "    ")
            mp.add_run("Status: ").bold = True;   mp.add_run(f"{a.get('status') or 'unanswered'}    ")
            mp.add_run("Owner: ").bold = True;    mp.add_run(f"{a.get('owner') or '—'}    ")
            mp.add_run("Follow-up: ").bold = True; mp.add_run(f"{a.get('follow_up') or 'none'}")
            doc.add_paragraph().add_run("Response").bold = True
            doc.add_paragraph(a.get("response") or "(no response captured)")
            if (a.get("evidence") or "").strip():
                doc.add_paragraph().add_run("Evidence / notes").bold = True
                doc.add_paragraph(a["evidence"])
            if (a.get("follow_up_note") or "").strip():
                doc.add_paragraph().add_run("Follow-up note").bold = True
                doc.add_paragraph(a["follow_up_note"])
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf


def _render_summary_docx(session, answers, summary):
    domain = session["domain"]; d = DOMAINS[domain]
    doc = Document(); _docx_styled(doc)
    label_map = {m["value"]: m["label"] for m in MODEL["maturity_scale"]}
    doc.add_heading(f"{d['title']} — Client Summary", level=0)
    p = doc.add_paragraph()
    p.add_run("Customer: ").bold = True; p.add_run((session.get("customer") or "—") + "\n")
    p.add_run("Domain: ").bold = True;   p.add_run(d.get("label", domain) + "\n")
    p.add_run("Date: ").bold = True;     p.add_run(datetime.now().strftime("%Y-%m-%d"))
    doc.add_heading("Maturity Snapshot", level=1)
    p2 = doc.add_paragraph(); p2.add_run("Overall: ").bold = True
    p2.add_run(f"{summary['overall_avg']}/5\n")
    table = doc.add_table(rows=1, cols=3); table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Layer"; hdr[1].text = "Title"; hdr[2].text = "Avg / 5"
    for row in summary["by_layer"]:
        r = table.add_row().cells
        r[0].text = f"L{row['layer_id']}"; r[1].text = row["title"]
        r[2].text = "—" if row["avg_score"] is None else str(row["avg_score"])
    doc.add_heading("Per-Question Maturity", level=1)
    qt = doc.add_table(rows=1, cols=4); qt.style = "Light Grid Accent 1"
    h = qt.rows[0].cells
    h[0].text = "Q"; h[1].text = "Layer"; h[2].text = "Question"; h[3].text = "Maturity"
    for row in summary["heatmap"]:
        r = qt.add_row().cells
        r[0].text = f"Q{row['question_id']}"; r[1].text = f"L{row['layer_id']}"
        r[2].text = row["primary"][:200]; r[3].text = row["maturity_label"] or "—"
    doc.add_heading("Gaps & Open Follow-ups", level=1)
    gap_rows = [h for h in summary["heatmap"] if (h["score"] is not None and h["score"] <= 2) or h["follow_up"] == "open"]
    if not gap_rows:
        doc.add_paragraph("No gaps or open follow-ups recorded.")
    else:
        gt = doc.add_table(rows=1, cols=3); gt.style = "Light Grid Accent 1"
        gh = gt.rows[0].cells
        gh[0].text = "Q"; gh[1].text = "Question"; gh[2].text = "Reason"
        for row in gap_rows:
            r = gt.add_row().cells
            r[0].text = f"Q{row['question_id']}"; r[1].text = row["primary"][:200]
            reasons = []
            if row["score"] is not None and row["score"] <= 2:
                reasons.append(f"Maturity {row['maturity_label']} ({row['score']}/5)")
            if row["follow_up"] == "open":
                reasons.append("Open follow-up")
            r[2].text = "; ".join(reasons)
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf

# ────────────────────────────── launcher ─────────────────────────────────────

def _pick_port(preferred: int = 8765) -> int:
    for port in [preferred, *range(8766, 8800)]:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("127.0.0.1", port)) != 0:
                return port
    return preferred

def _open_browser(port: int) -> None:
    def go():
        time.sleep(1.0)
        try: webbrowser.open_new(f"http://127.0.0.1:{port}/")
        except Exception: pass
    threading.Thread(target=go, daemon=True).start()

if __name__ == "__main__":
    port = _pick_port(int(os.environ.get("PORT", "8765")))
    print(f"\nInfrastructure Assessment  →  http://127.0.0.1:{port}/\n(Default admin: admin / Alpha#2026 — reset on first login)\n", file=sys.stderr)
    if "--no-browser" not in sys.argv:
        _open_browser(port)
    uvicorn.run(app, host="127.0.0.1", port=port, log_level="warning")
