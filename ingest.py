"""
ingest.py — Parse / author all eight domain questionnaires into a single canonical seed.

Domains produced (in display order):
  backup, storage, service_mgmt, identity, network, security_risk, compute, cloud_aws

- backup is parsed verbatim from src/Backup Discussion.docx and src/<Excel>/Backup tab.
- storage and the six new domains are authored in storage_content.py + domain_content.py,
  distilled from the corresponding Excel tabs. A parallel <Domain> Discussion.docx is
  generated under src/ for each authored domain so facilitators always have a Word artifact.

Outputs:
  - data/seed.json
  - src/<Domain> Discussion.docx (one per authored domain)
"""

from __future__ import annotations
import json, re, sys
from pathlib import Path
from typing import List, Dict, Any

from docx import Document
from docx.shared import Pt
import openpyxl

from storage_content import STORAGE_DOMAIN
from domain_content import DOMAINS as AUTHORED_DOMAINS

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "src"
DATA = ROOT / "data"
DATA.mkdir(exist_ok=True)

BACKUP_DOCX = SRC / "Backup Discussion.docx"
XLSX_PATH   = SRC / "DRAFT - GIC_Discovery_Questions_4.28.26.xlsx"

LAYER_PREFIXES = {f"Layer {i}:": i for i in range(1, 6)}
INTENT_VERBS   = ["Establish ", "Move from", "Probe ", "Test ", "Anchor "]

# Q8/Q9 in Backup Discussion.docx merge probes into a single paragraph. Split deterministically.
SUBQ_SPLITS = {
    8: [
        "Where are backup platform logs forwarded — enterprise SIEM, isolated cyber-recovery SOC, or only locally retained?",
        "What specific alerts fire for retention lock changes, immutability disables, mass deletes, or anomalous restore activity?",
        "What is the documented MTTD for backup tampering, and has it been measured rather than modelled?",
        "Are backup admin actions correlated with identity-side anomalies (privileged logons, new accounts, AD changes)?",
        "Are canary backups or honeypot artifacts deployed to detect intruder reconnaissance?",
    ],
    9: [
        "Which scanning engine runs against backups, on what cadence, and at what depth — signature, behavioural, ML?",
        "What happens when a backup is flagged — quarantined, alerted, or simply noted in a report?",
        "What does the clean room actually look like — isolated VLAN, separate vCenter, dedicated identity plane?",
        "How is the clean room itself kept clean and rebuilt between recovery cycles?",
        "Are integrity checks (checksum, hash, restore-verification) periodic, or only invoked at restore time?",
    ],
}


def parse_docx(path: Path) -> Dict[str, Any]:
    doc = Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    title = paras[0]
    facilitator_intro = ""
    if len(paras) > 1 and paras[1].lower().startswith("facilitator guidance"):
        facilitator_intro = paras[1].split(":", 1)[1].strip()

    layers: List[Dict[str, Any]] = []
    current_layer: Dict[str, Any] | None = None
    current_q: Dict[str, Any] | None = None
    primary_re = re.compile(r"^(\d{1,2})\.\s+(.*)")

    for text in paras[2:]:
        layer_hit = next((p for p in LAYER_PREFIXES if text.startswith(p)), None)
        if layer_hit:
            if current_q and current_layer:
                current_layer["questions"].append(current_q); current_q = None
            head = re.match(r"^Layer (\d):\s+(.*)", text)
            num, rest = int(head.group(1)), head.group(2)
            split_idx = -1
            for v in INTENT_VERBS:
                i = rest.find(v)
                if i != -1: split_idx = i; break
            if split_idx >= 0:
                layer_title, layer_intent = rest[:split_idx].strip(), rest[split_idx:].strip()
            else:
                layer_title, layer_intent = rest.strip(), ""
            current_layer = {"id": num, "title": layer_title, "intent": layer_intent, "questions": []}
            layers.append(current_layer)
            continue
        m = primary_re.match(text)
        if m and current_layer:
            if current_q: current_layer["questions"].append(current_q)
            current_q = {"id": int(m.group(1)), "layer_id": current_layer["id"],
                         "primary": m.group(2).strip(), "sub_questions": [], "facilitator_guidance": ""}
            continue
        if current_q is not None:
            current_q["sub_questions"].append(text)

    if current_q and current_layer:
        current_layer["questions"].append(current_q)
    for layer in layers:
        for q in layer["questions"]:
            if q["id"] in SUBQ_SPLITS:
                q["sub_questions"] = SUBQ_SPLITS[q["id"]]
    return {"title": title, "facilitator_intro": facilitator_intro, "layers": layers}


def parse_excel_sheet(sheet_name: str) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    if sheet_name not in wb.sheetnames:
        return {"taxonomy": {}, "granular_questions": []}
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h}
    discussion_col = "Discussion Questions (Broad Conversational)"
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row): continue
        rec = {
            "subdomain":  ((row[idx["Subdomain"]] or "") if idx.get("Subdomain") is not None else "").strip(),
            "topic":      ((row[idx["Topic"]] or "") if idx.get("Topic") is not None else "").strip(),
            "question":   ((row[idx["Question"]] or "") if idx.get("Question") is not None else "").strip(),
            "discussion": ((row[idx[discussion_col]] or "") if discussion_col in idx else "").strip(),
        }
        if rec["question"]: rows.append(rec)
    taxonomy: Dict[str, List[str]] = {}
    for r in rows:
        taxonomy.setdefault(r["subdomain"], [])
        if r["topic"] and r["topic"] not in taxonomy[r["subdomain"]]:
            taxonomy[r["subdomain"]].append(r["topic"])
    return {"taxonomy": taxonomy, "granular_questions": rows}


def write_discussion_docx(domain: Dict[str, Any], path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(11)
    doc.add_heading(domain["title"], level=1)
    doc.add_paragraph(f"Facilitator Guidance: {domain['facilitator_intro']}")
    for layer in domain["layers"]:
        doc.add_paragraph(f"Layer {layer['id']}: {layer['title']} {layer['intent']}")
        for q in layer["questions"]:
            doc.add_paragraph(f"{q['id']}. {q['primary']}")
            for sq in q["sub_questions"]: doc.add_paragraph(sq)
    doc.save(path)


MATURITY_SCALE = [
    {"value": "initial",    "label": "Initial",    "score": 1, "definition": "Ad-hoc, undocumented, dependent on individuals."},
    {"value": "developing", "label": "Developing", "score": 2, "definition": "Defined in some areas, inconsistent execution."},
    {"value": "defined",    "label": "Defined",    "score": 3, "definition": "Documented, standardised, repeatable across the estate."},
    {"value": "managed",    "label": "Managed",    "score": 4, "definition": "Measured, monitored, periodically tested against objectives."},
    {"value": "optimized",  "label": "Optimized",  "score": 5, "definition": "Continuously improved, evidence-led, validated under adverse conditions."},
]

DOMAIN_ORDER = ["backup", "storage", "service_mgmt", "identity", "network", "security_risk", "compute", "cloud_aws"]


def main() -> None:
    domains: Dict[str, Any] = {}

    # backup — from docx + excel
    backup = parse_docx(BACKUP_DOCX)
    backup_ex = parse_excel_sheet("Backup")
    domains["backup"] = {
        "key": "backup", "label": "Backup & IRE",
        "title": backup["title"], "facilitator_intro": backup["facilitator_intro"],
        "layers": backup["layers"],
        "taxonomy": backup_ex["taxonomy"], "granular_questions": backup_ex["granular_questions"],
    }

    # storage — authored + excel
    write_discussion_docx(STORAGE_DOMAIN, SRC / "Storage Discussion.docx")
    storage_ex = parse_excel_sheet("Storage")
    domains["storage"] = {
        "key": "storage", "label": "Storage Resilience",
        "title": STORAGE_DOMAIN["title"], "facilitator_intro": STORAGE_DOMAIN["facilitator_intro"],
        "layers": STORAGE_DOMAIN["layers"],
        "taxonomy": storage_ex["taxonomy"], "granular_questions": storage_ex["granular_questions"],
    }

    # other six authored domains
    for key, meta in AUTHORED_DOMAINS.items():
        d = meta["domain"]
        write_discussion_docx(d, SRC / f"{meta['label']} Discussion.docx")
        ex = parse_excel_sheet(meta["excel_sheet"])
        domains[key] = {
            "key": key, "label": meta["label"],
            "title": d["title"], "facilitator_intro": d["facilitator_intro"],
            "layers": d["layers"],
            "taxonomy": ex["taxonomy"], "granular_questions": ex["granular_questions"],
        }

    seed = {
        "domains": domains,
        "domain_order": DOMAIN_ORDER,
        "maturity_scale": MATURITY_SCALE,
        "follow_up_states": ["none", "open", "closed"],
        "answer_states":   ["unanswered", "in_progress", "complete"],
    }
    (DATA / "seed.json").write_text(json.dumps(seed, indent=2, ensure_ascii=False))

    for key in DOMAIN_ORDER:
        d = domains[key]
        nq = sum(len(l["questions"]) for l in d["layers"])
        nsub = sum(len(q["sub_questions"]) for l in d["layers"] for q in l["questions"])
        print(f"  {key:>14}: {d['label']:<26} {len(d['layers'])} layers, {nq:>2} Qs, {nsub:>3} subqs, {len(d['granular_questions']):>3} excel rows")
    print(f"\nseed.json → {DATA / 'seed.json'}")


if __name__ == "__main__":
    main()
